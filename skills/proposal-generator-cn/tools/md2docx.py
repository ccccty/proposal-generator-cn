# -*- coding: utf-8 -*-
"""Markdown -> Word (.docx) converter for Chinese research proposals.

特点
----
* 公式：$$...$$ 独立行、$...$ 行内，转 Word 原生 OMML（不依赖 MathType）。
* 图表：图片居中，自动题注防重复（alt 文字与下方 "图1-1 ..." 不再叠加）。
* 排版：字体/字号/页边距/行距/缩进/图宽全部可通过 --config 覆盖，
  默认值只是"可用的起点"，最终以本单位转发的模板为准。
* 自检：遇到不支持的 LaTeX 命令、缺失图片会汇总告警；--strict 下直接失败，
  避免字面反斜杠或空缺图混进交付件。

用法
----
    python md2docx.py 正文.md 正文.docx
    python md2docx.py 正文.md 正文.docx --config format.json --strict
    python md2docx.py --print-default-config > format.json
"""
import argparse
import json
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from latex_omml import (  # noqa: E402
    EQ_INLINE,
    extract_equations,
    latex_to_omml_inline,
    latex_to_omml_para,
    unknown_commands,
)

DEFAULT_FORMAT = {
    "fonts": {
        "body": "宋体",
        "heading": "黑体",
        "subheading": "楷体",
        "table": "宋体",
        "caption": "宋体",
        "title": "黑体",
        "latin": "Times New Roman",
    },
    "sizes": {
        "body": 12,
        "h1": 14,
        "h2": 12,
        "h3": 12,
        "table": 10.5,
        "caption": 10.5,
        "title": 18,
        "toc_title": 16,
    },
    "margins_cm": {"top": 2.5, "bottom": 2.5, "left": 2.8, "right": 2.8},
    "line_spacing": 1.5,
    "first_line_indent_chars": 2,
    "image_width_in": 5.5,
    "toc_levels": 3,
    "page_break_token": "<<PAGEBREAK>>",
}

BOLD_SEP = re.compile(r"(\*\*[^*]+\*\*)")
HR_LINE = re.compile(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$")
CAPTION_LINE = re.compile(r"^(?:图|表)\s?\d")
CODE_SPAN = re.compile("\x60([^\x60]*)\x60")

WARNINGS = []
STATS = {"paragraphs": 0, "tables": 0, "images": 0, "equations": 0, "headings": 0}


def warn(kind, message):
    WARNINGS.append((kind, message))


def _merge_format(cfg):
    merged = {k: (v.copy() if isinstance(v, dict) else v) for k, v in DEFAULT_FORMAT.items()}
    for k, v in (cfg or {}).items():
        if isinstance(v, dict) and isinstance(merged.get(k), dict):
            merged[k].update(v)
        else:
            merged[k] = v
    return merged


def load_format(path):
    if not path:
        return dict(DEFAULT_FORMAT)
    with open(path, encoding="utf-8-sig") as fh:
        return _merge_format(json.load(fh))


class Builder:
    def __init__(self, doc, fmt, base_dir):
        self.doc = doc
        self.fmt = fmt
        self.base_dir = base_dir
        self.latin = fmt["fonts"].get("latin", "Times New Roman")

    # ---- primitives -------------------------------------------------
    def set_run(self, run, cn_font, size, bold=False):
        run.font.name = self.latin
        run.font.size = Pt(size)
        run.bold = bold
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), cn_font)
        rfonts.set(qn("w:ascii"), self.latin)
        rfonts.set(qn("w:hAnsi"), self.latin)

    def _apply_spacing(self, p):
        pf = p.paragraph_format
        ls = self.fmt.get("line_spacing")
        if ls:
            pf.line_spacing = float(ls)
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE

    def _indent(self, p, chars):
        if not chars:
            return
        ind = OxmlElement("w:ind")
        ind.set(qn("w:firstLineChars"), str(int(chars * 100)))
        ind.set(qn("w:firstLine"), str(int(chars * 240)))
        p._p.get_or_add_pPr().append(ind)

    def _outline(self, p, lvl):
        pPr = p._p.get_or_add_pPr()
        o = OxmlElement("w:outlineLvl")
        o.set(qn("w:val"), str(lvl))
        pPr.append(o)

    def para(self, text, cn=None, size=None, bold=False, align=None,
             indent_chars=None, space_before=0, space_after=0, outline=None):
        f = self.fmt["fonts"]
        s = self.fmt["sizes"]
        cn = cn or f["body"]
        size = size or s["body"]
        if indent_chars is None:
            indent_chars = self.fmt.get("first_line_indent_chars", 2)
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        self._apply_spacing(p)
        if space_before:
            pf.space_before = Pt(space_before)
        if space_after:
            pf.space_after = Pt(space_after)
        if align is not None:
            pf.alignment = align
        self._indent(p, indent_chars)
        if outline is not None:
            self._outline(p, outline)
        self._emit_rich(p, text, cn, size, bold)
        STATS["paragraphs"] += 1
        return p

    def add_toc(self, levels=None):
        levels = levels or self.fmt.get("toc_levels", 3)
        p = self.doc.add_paragraph()
        self._apply_spacing(p)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), 'TOC \\o "1-%d" \\h \\z \\u' % levels)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "（在此处按 F9 或右键“更新域”生成目录）"
        r.append(t)
        fld.append(r)
        p._p.append(fld)
        return p

    def add_table(self, rows):
        if not rows:
            return
        ncol = max(len(r) for r in rows)
        t = self.doc.add_table(rows=0, cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        f = self.fmt["fonts"]
        size = self.fmt["sizes"]["table"]
        for i, row in enumerate(rows):
            cells = t.add_row().cells
            for j in range(ncol):
                txt = row[j] if j < len(row) else ""
                cell = cells[j]
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                if i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._emit_rich(p, txt, f["table"], size, bold=(i == 0))
        STATS["tables"] += 1
        return t

    def _insert_omml(self, p, omml_xml, display=False):
        from lxml import etree
        try:
            node = etree.fromstring(omml_xml.encode("utf-8"))
            if display:
                p._p.append(node)
            else:
                # m:oMath 是 w:p 的直接子元素，必须与 w:r 平级插入；
                # 放进 w:r 内部虽然 Word 可能容忍，但会给后续编辑带来结构问题。
                p._p.append(node)
            STATS["equations"] += 1
        except Exception as exc:  # pragma: no cover - defensive
            warn("equation", "公式渲染失败：%s" % str(exc)[:60])
            p.add_run("[公式解析失败]")

    def _inline_math(self, p, expr):
        bad = unknown_commands(expr)
        if bad:
            warn("latex", "不支持的公式命令 %s（原文：%s）" % ("、".join(bad), expr[:60]))
        self._insert_omml(p, latex_to_omml_inline(expr), display=False)

    def _emit_rich(self, p, text, cn, size, bold=False):
        """把一段文字就地写进段落：**加粗** 与 $行内公式$ 都不拆段、不换行。

        行内公式必须落在它所在的句子里——「其中 $G_{geo}$ 为由…，$T_{man}$ 为…」
        整句是一个段落，公式以 OMML 对象嵌在文字之间，不能各自独立成行。
        """
        for seg in BOLD_SEP.split(text):
            if not seg:
                continue
            b = bold
            if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
                seg = seg[2:-2]
                b = True
            pos = 0
            for m in EQ_INLINE.finditer(seg):
                if m.start() > pos:
                    self.set_run(p.add_run(seg[pos:m.start()]), cn, size, b)
                self._inline_math(p, m.group(1).strip())
                pos = m.end()
            if pos < len(seg):
                self.set_run(p.add_run(seg[pos:]), cn, size, b)

    # ---- block handlers --------------------------------------------
    def heading(self, level, text):
        f = self.fmt["fonts"]
        s = self.fmt["sizes"]
        STATS["headings"] += 1
        if level == 1:
            self.para(text, f["heading"], s["h1"], align=WD_ALIGN_PARAGRAPH.CENTER,
                      indent_chars=0, space_before=12, space_after=8, outline=0)
        elif level == 2:
            self.para(text, f["heading"], s["h1"], indent_chars=0,
                      space_before=14, space_after=6, outline=1)
        elif level == 3:
            self.para(text, f["subheading"], s["h2"], indent_chars=2,
                      space_before=8, space_after=4, outline=2)
        elif level == 4:
            self.para(text, f["heading"], s["h3"], bold=True, indent_chars=2,
                      space_before=5, outline=3)
        elif level == 5:
            self.para(text, f["body"], s["h3"], bold=True, indent_chars=2,
                      space_before=3, outline=4)
        else:
            self.para(text, f["body"], s["body"], bold=True, indent_chars=2)

    def image(self, alt, rel, next_line):
        caption = (alt or "").strip()
        img_path = rel if os.path.isabs(rel) else os.path.join(self.base_dir or ".", rel)
        if not os.path.isfile(img_path):
            warn("image", "图片缺失：%s" % rel)
            self.para("[图片缺失: %s]" % rel)
            return
        p = self.doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._apply_spacing(p)
        p.add_run().add_picture(img_path, width=Inches(self.fmt.get("image_width_in", 5.5)))
        STATS["images"] += 1
        # 题注防重复：alt 本身就是 "图1-1 ..." 或下一行紧接同一题注时只保留后者
        nxt = (next_line or "").strip()
        if caption and not CAPTION_LINE.match(caption) and nxt != caption and not CAPTION_LINE.match(nxt):
            self.para(caption, self.fmt["fonts"]["caption"], self.fmt["sizes"]["caption"],
                      align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0,
                      space_before=2, space_after=6)

    def caption(self, text):
        self.para(text, self.fmt["fonts"]["caption"], self.fmt["sizes"]["caption"],
                  align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0,
                  space_before=2, space_after=6)

    def equation_display(self, expr):
        bad = unknown_commands(expr)
        if bad:
            warn("latex", "不支持的公式命令 %s（原文：%s）" % ("、".join(bad), expr[:60]))
        p = self.para("", align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0,
                      space_before=2, space_after=2)
        self._insert_omml(p, latex_to_omml_para(expr), display=True)

    def equation_inline(self, expr):
        p = self.para("", indent_chars=0)
        self._inline_math(p, expr)


def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        sep = True
        for c in cells:
            if c and not re.fullmatch(r":?-{2,}:?", c):
                sep = False
        if not sep:
            rows.append(cells)
        i += 1
    return rows, i


def _clean_text(s):
    s = CODE_SPAN.sub(r"\1", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)  # 正文里的链接只留文字
    return s


def _clean_inline(s):
    """剥离代码标记与链接文字，但 $...$ 内部原样保留。

    公式里可能出现 \\left[ \\right] 之类的中括号，若整段一起过链接正则，
    公式内容会被当作 markdown 链接吃掉。因此按公式切成片段，只清理公式外的文字。
    """
    s = s.replace("\ufeff", "")
    out, pos = [], 0
    for m in EQ_INLINE.finditer(s):
        out.append(_clean_text(s[pos:m.start()]))
        out.append(m.group(0))
        pos = m.end()
    out.append(_clean_text(s[pos:]))
    return "".join(out)


def build(md, out, title=None, base_dir=None, fmt=None, toc=True):
    fmt = _merge_format(fmt or {})
    doc = Document()
    sec = doc.sections[0]
    m = fmt["margins_cm"]
    sec.top_margin = Cm(m["top"])
    sec.bottom_margin = Cm(m["bottom"])
    sec.left_margin = Cm(m["left"])
    sec.right_margin = Cm(m["right"])
    st = doc.styles["Normal"]
    st.font.name = fmt["fonts"].get("latin", "Times New Roman")
    st.font.size = Pt(fmt["sizes"]["body"])
    builder = Builder(doc, fmt, base_dir)

    if title:
        builder.para(title, fmt["fonts"]["title"], fmt["sizes"]["title"],
                     align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0,
                     space_before=24, space_after=24)
        if toc:
            builder.para("目  录", fmt["fonts"]["heading"], fmt["sizes"]["toc_title"],
                         align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0,
                         space_before=12, space_after=6)
            builder.add_toc()
            doc.add_page_break()

    try:
        from lxml import etree  # noqa: F401
        has_lxml = True
    except ImportError:
        has_lxml = False
        warn("env", "缺少 lxml，公式将退化为文字（pip install lxml）")

    segs = extract_equations(md) if has_lxml else [("md", md)]
    # 只有独立行公式切成单独块；行内公式按原样拼回所在句子，交给 para() 就地渲染。
    # 这里先按原文拼接再统一按行切分——否则「其中 $a$ 为…，$b$ 为…」会被拆成
    # 文字/公式/文字三行，产出多个段落后公式各自独立成行。
    buf = []
    for typ, content in segs:
        if typ == "md":
            buf.append(content)
        elif typ == "display":
            buf.append("\n<<EQ_DISPLAY:" + content.replace("\\", "\\\\") + ">>\n")
        else:
            buf.append("$" + content + "$")
    lines = "".join(buf).split("\n")

    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s == fmt.get("page_break_token", "<<PAGEBREAK>>"):
            doc.add_page_break()
            i += 1
            continue
        if HR_LINE.match(s):
            i += 1
            continue
        if s.startswith("|"):
            rows, i = parse_table(lines, i)
            builder.add_table(rows)
            continue
        mh = re.match(r"^(#{1,6})\s+(.*)$", s)
        if mh:
            builder.heading(len(mh.group(1)), _clean_inline(mh.group(2).strip()))
            i += 1
            continue
        m_img = re.match(r"^!\[(.*?)\]\((.+?)\)", s)
        if m_img:
            nxt = ""
            for k in range(i + 1, min(i + 4, len(lines))):
                if lines[k].strip():
                    nxt = lines[k].strip()
                    break
            builder.image(m_img.group(1), m_img.group(2).strip(), nxt)
            i += 1
            continue
        if CAPTION_LINE.match(s):
            builder.caption(_clean_inline(s))
            i += 1
            continue
        if s.startswith(">"):
            builder.para(_clean_inline(s.lstrip("> ").strip()), fmt["fonts"]["subheading"])
            i += 1
            continue
        if s.startswith("<<EQ_DISPLAY:") and has_lxml:
            builder.equation_display(s[13:-2].replace("\\\\", "\\"))
            i += 1
            continue
        if s.startswith("<<EQ_INLINE:") and has_lxml:
            builder.equation_inline(s[12:-2].replace("\\\\", "\\"))
            i += 1
            continue
        builder.para(_clean_inline(s))
        i += 1

    doc.save(out)
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(
        description="Markdown -> Word（中文申报书排版，公式转原生 OMML）")
    ap.add_argument("input", nargs="?", help="输入 .md 文件")
    ap.add_argument("output", nargs="?", help="输出 .docx 文件")
    ap.add_argument("--config", help="排版配置文件（JSON），见 --print-default-config")
    ap.add_argument("--title", help="封面标题，默认取第一个一级标题")
    ap.add_argument("--strict", action="store_true", help="有告警即失败（退出码 1）")
    ap.add_argument("--no-toc", action="store_true", help="不生成目录域")
    ap.add_argument("--quiet", action="store_true", help="只输出结果路径")
    ap.add_argument("--print-default-config", action="store_true",
                    help="打印默认排版配置（JSON）供修改")
    args = ap.parse_args(argv)

    if args.print_default_config:
        print(json.dumps(DEFAULT_FORMAT, ensure_ascii=False, indent=2))
        return 0

    if not args.input or not args.output:
        ap.print_help()
        return 2
    if not os.path.isfile(args.input):
        print("输入文件不存在：%s" % args.input, file=sys.stderr)
        return 2

    fmt = load_format(args.config)
    with open(args.input, encoding="utf-8-sig") as fh:
        md = fh.read()

    title = args.title
    if title is None:
        m = re.search(r"^#\s+(.+)$", md, re.M)
        if m:
            title = m.group(1).strip()
            md = md[:m.start()] + md[m.end():]

    out = build(md, args.output, title=title,
                base_dir=os.path.dirname(os.path.abspath(args.input)),
                fmt=fmt, toc=not args.no_toc)

    if not args.quiet:
        print("已生成：%s（%.1f KB）" % (out, os.path.getsize(out) / 1024.0))
        print("统计：段落 %d / 表格 %d / 图片 %d / 公式 %d / 标题 %d" % (
            STATS["paragraphs"], STATS["tables"], STATS["images"],
            STATS["equations"], STATS["headings"]))
        if WARNINGS:
            print("告警 %d 条：" % len(WARNINGS))
            for kind, msg in WARNINGS:
                print("  [%s] %s" % (kind, msg))
        else:
            print("告警：无")

    if args.strict and WARNINGS:
        print("--strict：存在告警，按失败处理", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
